"""Build the RUF Mega Workbook (.xlsx). Run: python3 build_workbook.py [output_path]"""

import sys
from openpyxl import Workbook

import sheets_core
import sheets_artifacts
import sheets_ai


def build(path):
    wb = Workbook()
    wb.remove(wb.active)

    sheets_core.build_home(wb)
    sheets_core.build_ruf101(wb)

    sheets_artifacts.build_project_statement(wb)
    sheets_artifacts.build_team_list(wb)
    sheets_artifacts.build_individual_accountabilities(wb)
    sheets_artifacts.build_rap(wb)
    sheets_artifacts.build_actions(wb)
    sheets_artifacts.build_weekly_schedule(wb)
    sheets_artifacts.build_wam(wb)
    sheets_artifacts.build_issues(wb)
    sheets_artifacts.build_decisions(wb)
    sheets_artifacts.build_cost_of_late(wb)
    sheets_artifacts.build_accountability_matrix(wb)
    sheets_artifacts.build_opportunity_sheet(wb)

    sheets_core.build_dashboard(wb)

    sheets_ai.build_ai_framework(wb)
    sheets_ai.build_ai_schema(wb)
    sheets_ai.build_ai_transcripts(wb)
    sheets_ai.build_ai_email(wb)

    wb.properties.title = "Risk Up Front — Project Mega Workbook"
    wb.properties.creator = "RUF Toolkit"
    wb.properties.description = ("Templatized Risk Up Front (RUF) project workbook: all four RUF documents plus "
                                 "supporting artifacts, dashboard, and AI-agent context & ingestion prompts.")
    wb.save(path)
    print(f"Saved {path} with {len(wb.sheetnames)} sheets:")
    for s in wb.sheetnames:
        print("  -", s)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "../RUF-Mega-Workbook.xlsx"
    build(out)
